"""Kurs koordinatori uchun hisobot eksporti: Excel va PDF (dekanat yig'ilishi, rahbariyatga taqdim etish).

Bo'limlar: umumiy ko'rsatkichlar, muammoli talabalar, davomati past (chegaraga yetgan) talabalar, akademik qarzdorlar,
kontrakt va trimestr qarzdorlari. Excel'da har bir bo'lim — alohida varaq; jami summalar va «Umumiy» varaqdagi sonlar
formulalar bilan hisoblanadi (jadval tahrirlansa ham to'g'ri qoladi). PDF — bitta hujjat, jadvallar bilan.
"""
from __future__ import annotations

import io
from pathlib import Path

import absence
import status
from config import HOURS_PER_PAIR, UNIVERSITY_NAME
from database import db
from tenancy import course_title, current_course, scope_students
from utils import fmt_date, fmt_dt, fmt_money, fmt_num, group_key, now, now_iso, parse_course

SECTIONS = {
    "prob": "Muammoli talabalar",
    "att": "Davomati past talabalar",
    "acad": "Akademik qarzdorlar",
    "kontrakt": "Kontrakt qarzdorlari",
    "trimestr": "Trimestr qarzdorlari",
}


# ---------------------------------------------------------------- ma'lumotlar
async def collect(f: str = "") -> dict:
    """Hisobot uchun barcha ma'lumotlar (panel bilan bir xil manba); f — kurs yoki guruh filtri."""
    students = scope_students(await db.fetchall("SELECT * FROM students ORDER BY group_name, full_name"))
    if f:
        course = parse_course(f) if ("kurs" in f.lower() or f.strip().isdigit()) else None
        gk = group_key(f)
        students = [s for s in students if (course and s["course"] == course) or (not course and s["group_key"] == gk)]
    sts = await status.all_statuses(students)
    linked = {r["student_id"] for r in await db.fetchall(
        """SELECT DISTINCT ps.student_id FROM parent_students ps JOIN parents p ON p.tg_id = ps.parent_id
           WHERE p.active = 1""")}
    dl = await status.deadlines()

    def base(x):
        st = x["student"]
        return {"name": st["full_name"], "group": st.get("group_name") or "", "hemis": st.get("hemis_id") or "",
                "linked": st["id"] in linked}

    prob = sorted((x for x in sts if any(x["flags"].values())),
                  key=lambda x: (-sum(x["flags"].values()), x["student"]["full_name"]))
    data = {
        "meta": {"university": UNIVERSITY_NAME, "course": course_title(current_course()) or "", "filter": f,
                 "generated": now_iso(), "deadlines": dl},
        "total": len(sts), "linked": sum(1 for x in sts if x["student"]["id"] in linked),
        "prob": [{**base(x), "count": sum(x["flags"].values()),
                  "pairs": (x["attendance"]["counted"] if x["flags"]["att"] else None),
                  "action": absence.action_for(x["level"]) if x["flags"]["att"] else "",
                  "acad": ", ".join(d["subject"] for d in x["debts"]),
                  "kontrakt": x["pays"]["kontrakt"]["debt"] if x["flags"]["kontrakt"] else None,
                  "trimestr": x["pays"]["trimestr"]["debt"] if x["flags"]["trimestr"] else None} for x in prob],
        "att": [{**base(x), "hours": x["attendance"]["counted"], "percent": x["attendance"].get("percent"),
                 "action": absence.action_for(x["level"]),
                 "source": (f"HEMIS {fmt_date(x['attendance']['as_of'], False)}" if x["attendance"]["source"] == "hemis"
                            else "kunlik davomat")}
                for x in sorted((x for x in sts if x["flags"]["att"]), key=lambda x: -x["attendance"]["counted"])],
        "acad": [{**base(x), "n": len(x["debts"]),
                  "subjects": "; ".join(f"{d['subject']} ({fmt_num(d['score'])} → «2», {d['semester']}-sem.)"
                                        for d in x["debts"])}
                 for x in sorted((x for x in sts if x["flags"]["acad"]), key=lambda x: -len(x["debts"]))],
    }
    for kind in ("kontrakt", "trimestr"):
        data[kind] = [{**base(x), "contract": x["pays"][kind].get("contract"), "paid": x["pays"][kind].get("paid"),
                       "percent": x["pays"][kind].get("percent"), "debt": x["pays"][kind]["debt"],
                       "as_of": x["pays"][kind]["as_of"]}
                      for x in sorted((x for x in sts if x["flags"][kind]), key=lambda x: -x["pays"][kind]["debt"])]
    return data


def file_name(data: dict, ext: str, only: str | None = None) -> str:
    part = {"prob": "muammoli", "att": "davomat", "acad": "akademik", "kontrakt": "kontrakt",
            "trimestr": "trimestr"}.get(only, "kurs_holati")
    course = (current_course() or "kurs").replace("/", "_")
    tail = f"_{group_key(data['meta']['filter'])}" if data["meta"]["filter"] else ""
    return f"{part}_{course}{tail}_{now().strftime('%Y-%m-%d')}.{ext}"


def title_line(data: dict) -> str:
    m = data["meta"]
    return " · ".join(x for x in (m["university"], m["course"], m["filter"]) if x)


# ---------------------------------------------------------------- Excel
def build_xlsx(data: dict, only: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    head_fill = PatternFill("solid", start_color="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##0" so\'m"'
    wb = Workbook()
    wb.remove(wb.active)
    ranges: dict[str, tuple[str, int, int]] = {}  # bo'lim -> (varaq nomi, birinchi va oxirgi ma'lumot qatori)

    def sheet(key: str, cols: list[tuple[str, int]], rows: list[list], formats: dict[int, str] | None = None,
              totals: tuple[int, ...] = ()):
        name = SECTIONS[key][:31]
        ws = wb.create_sheet(name)
        ws["A1"] = SECTIONS[key]
        ws["A1"].font = Font(name=F, bold=True, size=14)
        ws["A2"] = f"{title_line(data)} · {fmt_dt(data['meta']['generated'])} holatiga"
        ws["A2"].font = Font(name=F, size=10, color="595959")
        h = 4
        for j, (title, width) in enumerate(cols, start=1):
            c = ws.cell(row=h, column=j, value=title)
            c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.row_dimensions[h].height = 32
        for i, row in enumerate(rows, start=h + 1):
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=F, size=10)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=j in (2, len(cols)))
                if formats and j in formats:
                    c.number_format = formats[j]
        first, last = h + 1, h + max(len(rows), 1)
        ranges[key] = (name, first, last)  # bo'sh ro'yxatda — bitta bo'sh qator (COUNTA = 0)
        if rows:
            ws.auto_filter.ref = f"A{h}:{get_column_letter(len(cols))}{last}"
            if totals:
                t = last + 1
                ws.cell(row=t, column=2, value="Jami").font = Font(name=F, bold=True, size=10)
                for j in totals:
                    col = get_column_letter(j)
                    c = ws.cell(row=t, column=j, value=f"=SUM({col}{first}:{col}{last})")
                    c.font = Font(name=F, bold=True, size=10)
                    c.number_format = (formats or {}).get(j, "General")
                    c.border = border
        else:  # izoh sanaladigan diapazondan tashqarida — «Umumiy» varaqdagi COUNTA 0 ko'rsatadi
            ws.cell(row=h + 3, column=2, value="Ro'yxat bo'sh").font = Font(name=F, italic=True, size=10)
        ws.freeze_panes = ws.cell(row=h + 1, column=3)
        ws.sheet_view.zoomScale = 100
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{h}:{h}"

    want = [only] if only else list(SECTIONS)
    if "prob" in want:
        sheet("prob", [("№", 5), ("F.I.Sh.", 34), ("Guruh", 11), ("HEMIS ID", 15), ("Muammolar soni", 10),
                       ("Sababsiz qoldirilgan, para", 12), ("Sababsiz qoldirilgan, soat", 12), ("Chora (chegara)", 26),
                       ("Akademik qarz (fanlar)", 30), ("Kontrakt qarzi", 15), ("Trimestr qarzi", 15),
                       ("Ota-ona botda", 10)],
              [[i, r["name"], r["group"], r["hemis"], r["count"],
                None if r["pairs"] is None else f"=G{4 + i}/{fmt_num(HOURS_PER_PAIR)}", r["pairs"], r["action"], r["acad"],
                r["kontrakt"], r["trimestr"], "ha" if r["linked"] else "yo'q"] for i, r in enumerate(data["prob"], 1)],
              {6: "0.#", 7: "0.#", 10: money, 11: money}, totals=(10, 11))
    if "att" in want:
        sheet("att", [("№", 5), ("F.I.Sh.", 34), ("Guruh", 11), ("HEMIS ID", 15), ("Sababsiz qoldirilgan, para", 12),
                      ("Sababsiz qoldirilgan, soat", 12), ("Davomat", 10), ("Chora (chegara)", 28), ("Manba", 16),
                      ("Ota-ona botda", 10)],
              [[i, r["name"], r["group"], r["hemis"], f"=F{4 + i}/{fmt_num(HOURS_PER_PAIR)}", r["hours"],
                None if r["percent"] is None else r["percent"] / 100, r["action"], r["source"],
                "ha" if r["linked"] else "yo'q"] for i, r in enumerate(data["att"], 1)],
              {5: "0.#", 6: "0.#", 7: "0%"})
    if "acad" in want:
        sheet("acad", [("№", 5), ("F.I.Sh.", 34), ("Guruh", 11), ("HEMIS ID", 15), ("Qarzdor fanlar soni", 10),
                       ("Fanlar (ball → baho, semestr)", 60), ("Ota-ona botda", 10)],
              [[i, r["name"], r["group"], r["hemis"], r["n"], r["subjects"], "ha" if r["linked"] else "yo'q"]
               for i, r in enumerate(data["acad"], 1)], totals=(5,))
    for kind in ("kontrakt", "trimestr"):
        if kind in want:
            dl = data["meta"]["deadlines"].get(kind)
            sheet(kind, [("№", 5), ("F.I.Sh.", 34), ("Guruh", 11), ("HEMIS ID", 15), ("Shartnoma summasi", 16),
                         ("To'langan", 16), ("To'langan, %", 10), ("Qarzdorlik", 16), ("Hisobot sanasi", 12),
                         ("To'lov muddati", 12), ("Ota-ona botda", 10)],
                  [[i, r["name"], r["group"], r["hemis"], r["contract"], r["paid"],
                    None if r["percent"] is None else r["percent"] / 100, r["debt"], fmt_date(r["as_of"], False),
                    fmt_date(dl, False) if dl else "—", "ha" if r["linked"] else "yo'q"]
                   for i, r in enumerate(data[kind], 1)],
                  {5: money, 6: money, 7: "0%", 8: money}, totals=(8,))

    if not only:  # «Umumiy» varaq — ko'rsatkichlar boshqa varaqlardan formulalar bilan
        ws = wb.create_sheet("Umumiy", 0)
        ws["A1"] = "Kurs holati — hisobot"
        ws["A1"].font = Font(name=F, bold=True, size=14)
        ws["A2"] = title_line(data)
        ws["A3"] = f"Tayyorlangan: {fmt_dt(data['meta']['generated'])}"
        for c in ("A2", "A3"):
            ws[c].font = Font(name=F, size=10, color="595959")

        def ref(key, col):
            name, first, last = ranges[key]
            return f"'{name}'!{col}{first}:{col}{last}"
        rows = [
            ("Jami talabalar", data["total"], None, "Bazadagi talabalar soni (bot hisobi)"),
            ("Ota-onasi botga ulangan talabalar", data["linked"], None, "Bot hisobi"),
            ("Muammoli talabalar (kamida 1 muammo)", f"=COUNTA({ref('prob', 'B')})", None, "«Muammoli talabalar» varag'idan"),
            ("3 va undan ortiq muammoli talabalar", f"=COUNTIF({ref('prob', 'E')},\">=3\")", None, "«Muammolar soni» ≥ 3"),
            ("Davomati past (chegaraga yetgan)", f"=COUNTA({ref('att', 'B')})", None,
             f"Sababsiz qoldirilgan ≥ {fmt_num(absence.LEVELS[0] / HOURS_PER_PAIR)} para ({absence.LEVELS[0]} soat)"),
            ("Akademik qarzdorlar", f"=COUNTA({ref('acad', 'B')})", None, "Kamida bitta fandan «2» (0–59 ball)"),
            ("Kontrakt qarzdorlari", f"=COUNTA({ref('kontrakt', 'B')})", f"=SUM({ref('kontrakt', 'H')})", "Oxirgi buxgalteriya hisoboti"),
            ("Trimestr qarzdorlari", f"=COUNTA({ref('trimestr', 'B')})", f"=SUM({ref('trimestr', 'H')})", "Oxirgi buxgalteriya hisoboti"),
        ]
        for j, (t, w) in enumerate((("Ko'rsatkich", 40), ("Talabalar", 12), ("Summa", 18), ("Izoh", 48)), start=1):
            c = ws.cell(row=5, column=j, value=t)
            c.font = Font(name=F, bold=True, color="FFFFFF", size=10)
            c.fill = head_fill
            c.border = border
            ws.column_dimensions[get_column_letter(j)].width = w
        for i, (label, val, total, note) in enumerate(rows, start=6):
            for j, v in enumerate((label, val, total, note), start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = Font(name=F, size=10, color="595959" if j == 4 else "000000")
                c.border = border
            ws.cell(row=i, column=3).number_format = money
        ws.cell(row=15, column=1, value="Ota-onasi ulangan talabalar ulushi").font = Font(name=F, size=10)
        c = ws.cell(row=15, column=2, value="=IF(B6=0,0,B7/B6)")
        c.number_format = "0%"
        c.font = Font(name=F, size=10)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
_FONT_DIRS = ("/usr/share/fonts/truetype/dejavu", "/usr/share/fonts/dejavu", "C:/Windows/Fonts")


def _fonts() -> tuple[str, str]:
    """Kirill va o'zbek harflari (ʻ ‘) uchun DejaVu; topilmasa — Helvetica (faqat lotin)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "DV" in pdfmetrics.getRegisteredFontNames():
        return "DV", "DVB"
    dirs = list(_FONT_DIRS)
    try:  # matplotlib o'zi bilan DejaVu shriftini olib keladi — serverga alohida shrift o'rnatish shart emas
        import matplotlib
        dirs.append(str(Path(matplotlib.get_data_path()) / "fonts" / "ttf"))
    except ImportError:
        pass
    for d in dirs:
        reg, bold = Path(d) / "DejaVuSans.ttf", Path(d) / "DejaVuSans-Bold.ttf"
        if reg.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("DV", str(reg)))
            pdfmetrics.registerFont(TTFont("DVB", str(bold)))
            return "DV", "DVB"
    return "Helvetica", "Helvetica-Bold"


def build_pdf(data: dict, only: str | None = None) -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    reg, bold = _fonts()
    st_title = ParagraphStyle("t", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    st_sub = ParagraphStyle("s", fontName=reg, fontSize=9, leading=12, textColor=colors.HexColor("#595959"))
    st_h = ParagraphStyle("h", fontName=bold, fontSize=12, leading=15, spaceBefore=10, spaceAfter=5,
                          textColor=colors.HexColor("#1F4E78"), keepWithNext=1)  # sarlavha jadvalidan ajralmaydi
    st_cell = ParagraphStyle("c", fontName=reg, fontSize=8, leading=10)
    st_head = ParagraphStyle("hc", fontName=bold, fontSize=8, leading=10, textColor=colors.white)
    blue = colors.HexColor("#1F4E78")

    def P(v, style=st_cell):
        return Paragraph(escape("" if v is None else str(v)), style)

    def table(headers, rows, widths):
        body = [[P(h, st_head) for h in headers]] + [[P(v) for v in r] for r in rows]
        t = Table(body, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), blue),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FA")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        return t

    def money(v):
        return fmt_money(v) if v is not None else ""

    def pct(v):
        return f"{round(v)}%" if v is not None else ""

    story = [Paragraph(escape("Kurs holati — hisobot" if not only else SECTIONS[only]), st_title),
             Paragraph(escape(f"{title_line(data)} · tayyorlangan: {fmt_dt(data['meta']['generated'])}"), st_sub),
             Spacer(1, 6)]
    if not only:
        k, t = data["kontrakt"], data["trimestr"]
        kpi = [["Jami talabalar", data["total"], ""],
               ["Ota-onasi botga ulangan talabalar", data["linked"],
                f"{round(100 * data['linked'] / data['total'])}%" if data["total"] else ""],
               ["Muammoli talabalar (kamida 1 muammo)", len(data["prob"]), ""],
               ["3 va undan ortiq muammoli talabalar", sum(1 for r in data["prob"] if r["count"] >= 3), ""],
               [f"Davomati past (≥ {fmt_num(absence.LEVELS[0] / HOURS_PER_PAIR)} para / {absence.LEVELS[0]} soat sababsiz)",
                len(data["att"]), ""],
               ["Akademik qarzdorlar", len(data["acad"]), ""],
               ["Kontrakt qarzdorlari", len(k), fmt_money(sum(r["debt"] for r in k))],
               ["Trimestr qarzdorlari", len(t), fmt_money(sum(r["debt"] for r in t))]]
        story += [Paragraph("Umumiy ko'rsatkichlar", st_h), table(["Ko'rsatkich", "Talabalar", "Summa / ulush"], kpi, [120, 30, 50])]
    want = [only] if only else list(SECTIONS)
    specs = {
        "prob": (["№", "F.I.Sh.", "Guruh", "Muammo", "Sababsiz qoldirilgan", "Chora", "Akademik qarz", "Kontrakt",
                  "Trimestr"], [8, 50, 22, 20, 32, 36, 50, 27, 27],
                 lambda i, r: [i, r["name"], r["group"], r["count"],
                               "" if r["pairs"] is None else f"{fmt_num(r['pairs'] / HOURS_PER_PAIR)} para ({fmt_num(r['pairs'])} soat)",
                               r["action"], r["acad"], money(r["kontrakt"]), money(r["trimestr"])]),
        "att": (["№", "F.I.Sh.", "Guruh", "HEMIS ID", "Sababsiz qoldirilgan", "Davomat", "Chora (chegara)", "Manba",
                 "Ota-ona botda"], [8, 56, 22, 28, 34, 20, 50, 32, 22],
                lambda i, r: [i, r["name"], r["group"], r["hemis"],
                              f"{fmt_num(r['hours'] / HOURS_PER_PAIR)} para ({fmt_num(r['hours'])} soat)",
                              pct(r["percent"]), r["action"], r["source"], "ha" if r["linked"] else "yo'q"]),
        "acad": (["№", "F.I.Sh.", "Guruh", "Fanlar soni", "Fanlar (ball → baho, semestr)", "Ota-ona botda"],
                 [8, 56, 22, 18, 146, 22], lambda i, r: [i, r["name"], r["group"], r["n"], r["subjects"],
                                                         "ha" if r["linked"] else "yo'q"]),
    }
    for kind in ("kontrakt", "trimestr"):
        dl = data["meta"]["deadlines"].get(kind)
        specs[kind] = (["№", "F.I.Sh.", "Guruh", "Shartnoma", "To'langan", "%", "Qarzdorlik", "Hisobot sanasi", "Muddat"],
                       [8, 62, 22, 34, 34, 16, 34, 32, 30],
                       lambda i, r, dl=dl: [i, r["name"], r["group"], money(r["contract"]), money(r["paid"]),
                                            pct(r["percent"]), money(r["debt"]), fmt_date(r["as_of"], False),
                                            fmt_date(dl, False) if dl else "—"])
    for key in want:
        headers, widths, fn = specs[key]
        rows = [fn(i, r) for i, r in enumerate(data[key], 1)]
        story.append(Paragraph(f"{escape(SECTIONS[key])} — {len(rows)} ta", st_h))
        if key in ("kontrakt", "trimestr") and rows:
            rows.append(["", "Jami", "", "", "", "", money(sum(r["debt"] for r in data[key])), "", ""])
        story.append(table(headers, rows, widths) if rows else Paragraph("Ro'yxat bo'sh.", st_cell))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(reg, 7)
        canvas.setFillColor(colors.HexColor("#808080"))
        canvas.drawString(12 * mm, 7 * mm, f"Ota-onalar davomat boti · {title_line(data)}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"{doc.page}-bet")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title="Kurs holati — hisobot", author="Ota-onalar davomat boti")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
