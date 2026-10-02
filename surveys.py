"""Ota-onalar o'rtasida so'rovnomalar.

Kurs koordinatori (yoki super-admin) so'rovnoma tuzadi: savollar — bitta javob, bir nechta javob, 1–5 baho yoki
erkin matn. So'rovnoma kursning barcha ota-onalariga yoki tanlangan guruhlarga yuboriladi (guruhlari biriktirilgan
koordinator — faqat o'z guruhlariga). Ota-ona ilovada javob beradi; bitta ota-ona — bitta javob (yopilguncha
o'zgartirishi mumkin). Natijalar — har bir savol bo'yicha sonlar va foizlar, erkin javoblar ro'yxati, Excel.

Ma'lumotlar kurs bazasida (data/<kurs>/bot.db): surveys, survey_questions, survey_responses, survey_answers.
"""
from __future__ import annotations

import io
import json
import logging

from database import db
from utils import group_key, now_iso, today

log = logging.getLogger("surveys")

KINDS = ("single", "multi", "scale", "text")
KIND_TITLES = {"single": "Bitta javob", "multi": "Bir nechta javob", "scale": "Baho (1–5)", "text": "Erkin javob"}
MAX_QUESTIONS, MAX_OPTIONS, MAX_TEXT = 30, 12, 2000


def _groups(s: dict) -> list[str]:
    return [g for g in (s.get("target_groups") or "").split(",") if g]


def is_open(s: dict) -> bool:
    return not s["closed"] and (not s.get("closes_at") or s["closes_at"] >= today().isoformat())


# ---------------------------------------------------------------- tuzish (kurs koordinatori)
def clean_questions(raw) -> list[dict]:
    """Ilovadan kelgan savollarni tekshiradi. Xato bo'lsa ValueError (sababi bilan)."""
    if not isinstance(raw, list) or not raw:
        raise ValueError("Kamida bitta savol kerak")
    out = []
    for i, q in enumerate(raw[:MAX_QUESTIONS], 1):
        kind = q.get("kind") if isinstance(q, dict) else None
        text = str((q or {}).get("text") or "").strip()[:500]
        if kind not in KINDS or not text:
            raise ValueError(f"{i}-savol to'liq emas")
        opts = []
        if kind in ("single", "multi"):
            opts = [str(o).strip()[:200] for o in (q.get("options") or []) if str(o).strip()][:MAX_OPTIONS]
            if len(opts) < 2:
                raise ValueError(f"{i}-savolda kamida ikkita javob varianti bo'lsin")
        out.append({"kind": kind, "text": text, "options": opts, "required": 1 if q.get("required", True) else 0})
    return out


async def create(title: str, description: str, questions: list[dict], groups: list[str] | None, by: int,
                 closes_at: str | None = None, anonymous: bool = False) -> int:
    sid = await db.execute(
        "INSERT INTO surveys (title, description, target_groups, anonymous, created_by, created_at, closes_at) "
        "VALUES (?, ?, ?, ?, ?, ?, ?)",
        (title.strip()[:200], (description or "").strip()[:2000] or None, ",".join(groups or []), int(anonymous), by,
         now_iso(), closes_at or None))
    for pos, q in enumerate(questions, 1):
        await db.execute("INSERT INTO survey_questions (survey_id, pos, kind, text, options, required) VALUES (?,?,?,?,?,?)",
                         (sid, pos, q["kind"], q["text"], json.dumps(q["options"], ensure_ascii=False), q["required"]))
    log.info("So'rovnoma yaratildi: #%s «%s», guruhlar: %s (%s)", sid, title, ",".join(groups or []) or "hammasi", by)
    return sid


async def get(sid: int) -> dict | None:
    s = await db.fetchone("SELECT * FROM surveys WHERE id = ?", (sid,))
    if not s:
        return None
    qs = await db.fetchall("SELECT * FROM survey_questions WHERE survey_id = ? ORDER BY pos", (sid,))
    return {**s, "groups": _groups(s), "questions": [{**q, "options": json.loads(q["options"] or "[]")} for q in qs]}


async def recipients(s: dict) -> list[int]:
    """Faol ota-onalar: so'rovnoma guruhlaridagi (yoki butun kursdagi) talabalarning ota-onalari."""
    return await db.parent_ids_for_groups(_groups(s) or None)


# ---------------------------------------------------------------- ota-ona
async def _child_for(parent_id: int, s: dict) -> dict | None:
    """Ota-onaning shu so'rovnomaga tegishli farzandi (guruhi mos). None — so'rovnoma unga tegishli emas."""
    gs = set(_groups(s))
    for c in await db.parent_children(parent_id):
        if not gs or (c.get("group_key") or "") in gs:
            return c
    return None


async def for_parent(parent_id: int) -> list[dict]:
    """Joriy kursdagi ota-onaga tegishli so'rovnomalar (ochiq va javob berilganlari), yangidan eskiga."""
    if not await db.get_parent(parent_id):
        return []
    out = []
    for s in await db.fetchall("SELECT * FROM surveys ORDER BY id DESC LIMIT 50"):
        if not await _child_for(parent_id, s):
            continue
        r = await db.fetchone("SELECT submitted_at FROM survey_responses WHERE survey_id = ? AND parent_id = ?",
                              (s["id"], parent_id))
        n = (await db.fetchone("SELECT COUNT(*) n FROM survey_questions WHERE survey_id = ?", (s["id"],)))["n"]
        if not is_open(s) and not r:
            continue  # yopilgan va javob berilmagan — ko'rsatilmaydi
        out.append({"id": s["id"], "title": s["title"], "description": s["description"], "questions": n,
                    "closes_at": s["closes_at"], "open": is_open(s), "answered": bool(r),
                    "submitted_at": r["submitted_at"] if r else None, "created_at": s["created_at"]})
    return out


async def my_answers(sid: int, parent_id: int) -> dict[int, str]:
    rows = await db.fetchall("SELECT question_id, value FROM survey_answers WHERE survey_id = ? AND parent_id = ?",
                             (sid, parent_id))
    return {r["question_id"]: r["value"] for r in rows}


async def submit(sid: int, parent_id: int, answers: dict) -> str | None:
    """Javobni tekshiradi va saqlaydi. Qaytaradi: xato kodi yoki None (muvaffaqiyat)."""
    s = await get(sid)
    if not s:
        return "not_found"
    if not is_open(s):
        return "closed"
    child = await _child_for(parent_id, s)
    if not child:
        return "forbidden"
    clean: dict[int, str] = {}
    for q in s["questions"]:
        v = answers.get(str(q["id"]), answers.get(q["id"]))
        if q["kind"] == "single":
            v = str(v) if v is not None and str(v) in q["options"] else None
        elif q["kind"] == "multi":
            v = [str(x) for x in (v or []) if str(x) in q["options"]] if isinstance(v, list) else []
            v = json.dumps(v, ensure_ascii=False) if v else None
        elif q["kind"] == "scale":
            v = str(v) if str(v) in ("1", "2", "3", "4", "5") else None
        else:
            v = str(v or "").strip()[:MAX_TEXT] or None
        if v is None:
            if q["required"]:
                return f"required:{q['id']}"
            continue
        clean[q["id"]] = v
    await db.execute("DELETE FROM survey_answers WHERE survey_id = ? AND parent_id = ?", (sid, parent_id))
    for qid, v in clean.items():
        await db.execute("INSERT INTO survey_answers (survey_id, parent_id, question_id, value) VALUES (?, ?, ?, ?)",
                         (sid, parent_id, qid, v))
    await db.execute(
        """INSERT INTO survey_responses (survey_id, parent_id, student_id, group_key, submitted_at) VALUES (?,?,?,?,?)
           ON CONFLICT(survey_id, parent_id) DO UPDATE SET student_id = excluded.student_id,
             group_key = excluded.group_key, submitted_at = excluded.submitted_at""",
        (sid, parent_id, child["id"], child.get("group_key") or "", now_iso()))
    return None


# ---------------------------------------------------------------- natijalar (kurs koordinatori)
def visible_to(s: dict, uid: int, scope: set[str] | None) -> bool:
    """Koordinator — o'zi tuzgan yoki o'z guruhlariga tegishli so'rovnomalar; super-admin/guruhsiz — hammasi."""
    if scope is None or s.get("created_by") == uid:
        return True
    gs = set(_groups(s))
    return bool(gs) and bool(gs & scope)


async def listing(uid: int, scope: set[str] | None) -> list[dict]:
    out = []
    for s in await db.fetchall("SELECT * FROM surveys ORDER BY id DESC"):
        if not visible_to(s, uid, scope):
            continue
        gs = _groups(s)
        reach = [g for g in gs if scope is None or g in scope] or (sorted(scope) if scope is not None else None)
        eligible = len(await db.parent_ids_for_groups(reach))
        answered = await _count_responses(s["id"], scope)
        names = await _group_names(gs)
        out.append({"id": s["id"], "title": s["title"], "created_at": s["created_at"], "closes_at": s["closes_at"],
                    "open": is_open(s), "groups": names, "answered": answered, "eligible": eligible,
                    "mine": s["created_by"] == uid})
    return out


async def _count_responses(sid: int, scope: set[str] | None) -> int:
    rows = await db.fetchall("SELECT group_key FROM survey_responses WHERE survey_id = ?", (sid,))
    return sum(1 for r in rows if scope is None or (r["group_key"] or "") in scope)


async def _group_names(keys: list[str]) -> list[str]:
    if not keys:
        return []
    names = {r["group_key"]: r["n"] for r in await db.fetchall(
        "SELECT group_key, MAX(group_name) n FROM students GROUP BY group_key")}
    return [names.get(k, k) for k in keys]


async def results(sid: int, scope: set[str] | None) -> dict | None:
    """Har bir savol bo'yicha natija (koordinator — faqat o'z guruhlari ota-onalarining javoblari)."""
    s = await get(sid)
    if not s:
        return None
    resp = [r for r in await db.fetchall(
        """SELECT r.*, s.full_name AS student, s.group_name, p.tg_name AS parent FROM survey_responses r
           LEFT JOIN students s ON s.id = r.student_id LEFT JOIN parents p ON p.tg_id = r.parent_id
           WHERE r.survey_id = ? ORDER BY r.submitted_at""", (sid,)) if scope is None or (r["group_key"] or "") in scope]
    who = {r["parent_id"]: r for r in resp}
    answers = [a for a in await db.fetchall("SELECT * FROM survey_answers WHERE survey_id = ?", (sid,))
               if a["parent_id"] in who]
    gs = s["groups"]
    reach = [g for g in gs if scope is None or g in scope] or (sorted(scope) if scope is not None else None)
    qs = []
    for q in s["questions"]:
        mine = [a for a in answers if a["question_id"] == q["id"]]
        item = {"id": q["id"], "kind": q["kind"], "text": q["text"], "answered": len(mine)}
        if q["kind"] in ("single", "multi", "scale"):
            opts = q["options"] if q["kind"] != "scale" else ["1", "2", "3", "4", "5"]
            counts = {o: 0 for o in opts}
            for a in mine:
                for v in (json.loads(a["value"]) if q["kind"] == "multi" else [a["value"]]):
                    if v in counts:
                        counts[v] += 1
            item["options"] = [{"label": o, "count": c} for o, c in counts.items()]
            if q["kind"] == "scale" and mine:
                item["average"] = round(sum(int(a["value"]) for a in mine) / len(mine), 2)
        else:
            item["texts"] = [{"text": a["value"],
                              **({} if s["anonymous"] else {"student": who[a["parent_id"]]["student"],
                                                           "group": who[a["parent_id"]]["group_name"]})}
                             for a in mine]
        qs.append(item)
    return {"id": s["id"], "title": s["title"], "description": s["description"], "open": is_open(s),
            "closes_at": s["closes_at"], "anonymous": bool(s["anonymous"]), "created_at": s["created_at"],
            "groups": await _group_names(gs), "answered": len(resp),
            "eligible": len(await db.parent_ids_for_groups(reach)), "questions": qs,
            "respondents": [] if s["anonymous"] else [
                {"student": r["student"], "group": r["group_name"], "parent": r["parent"], "at": r["submitted_at"]}
                for r in resp]}


async def close(sid: int) -> None:
    await db.execute("UPDATE surveys SET closed = 1 WHERE id = ?", (sid,))


async def delete(sid: int) -> None:
    await db.execute("DELETE FROM surveys WHERE id = ?", (sid,))


def build_xlsx(res: dict, raw: list[dict] | None = None) -> bytes:
    """Natijalar Excel'da: «Natijalar» (savollar bo'yicha sonlar) va «Javoblar» (har bir ota-ona — bitta qator)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Natijalar"
    ws["A1"] = res["title"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Javob berganlar: {res['answered']} / {res['eligible']}"
    row = 4
    for i, q in enumerate(res["questions"], 1):
        ws.cell(row=row, column=1, value=f"{i}. {q['text']}").font = Font(bold=True)
        row += 1
        if "options" in q:
            total = sum(o["count"] for o in q["options"]) or 1
            for o in q["options"]:
                ws.cell(row=row, column=1, value=o["label"])
                ws.cell(row=row, column=2, value=o["count"])
                c = ws.cell(row=row, column=3, value=o["count"] / total)
                c.number_format = "0%"
                row += 1
            if q.get("average") is not None:
                ws.cell(row=row, column=1, value="O'rtacha baho")
                ws.cell(row=row, column=2, value=q["average"])
                row += 1
        else:
            for tx in q.get("texts", []):
                c = ws.cell(row=row, column=1, value=tx["text"])
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if tx.get("student"):
                    ws.cell(row=row, column=2, value=f"{tx['student']} ({tx.get('group') or ''})")
                row += 1
        row += 1
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10
    if raw is not None:
        ws2 = wb.create_sheet("Javoblar")
        ws2.append(["Talaba", "Guruh", "Vaqt"] + [f"{i}. {q['text']}" for i, q in enumerate(res["questions"], 1)])
        for r in raw:
            ws2.append([r.get("student") or "", r.get("group") or "", r.get("at") or ""]
                       + [r["answers"].get(q["id"], "") for q in res["questions"]])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def raw_rows(sid: int, scope: set[str] | None) -> list[dict]:
    """Excel «Javoblar» varag'i uchun: har bir ota-onaning barcha javoblari (anonim so'rovnomada ism yo'q)."""
    s = await get(sid)
    out = []
    for r in await db.fetchall(
            """SELECT r.parent_id, r.group_key, r.submitted_at, s.full_name, s.group_name FROM survey_responses r
               LEFT JOIN students s ON s.id = r.student_id WHERE r.survey_id = ? ORDER BY r.submitted_at""", (sid,)):
        if scope is not None and (r["group_key"] or "") not in scope:
            continue
        ans = {}
        for a in await db.fetchall("SELECT question_id, value FROM survey_answers WHERE survey_id = ? AND parent_id = ?",
                                   (sid, r["parent_id"])):
            v = a["value"]
            if v.startswith("["):
                v = ", ".join(json.loads(v))
            ans[a["question_id"]] = v
        out.append({"student": None if s["anonymous"] else r["full_name"], "group": r["group_name"],
                    "at": r["submitted_at"][:16].replace("T", " "), "answers": ans})
    return out


def target_keys(groups_raw, scope: set[str] | None) -> list[str] | None:
    """Ilovadan kelgan guruh nomlari → kalitlar; koordinator — faqat o'z guruhlaridan (bo'sh — barcha guruhlari).
    Qaytaradi: kalitlar ro'yxati ([] — butun kurs) yoki None (koordinatorga tegishli guruh tanlanmadi)."""
    keys = [group_key(g) for g in (groups_raw or []) if group_key(g)]
    if scope is None:
        return keys
    keys = [k for k in keys if k in scope] if keys else sorted(scope)
    return keys or None
